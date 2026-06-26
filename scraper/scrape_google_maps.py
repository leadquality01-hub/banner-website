"""
Scraper Avocats Paris - Google Maps (via Playwright)
Google Maps n'a pas de Cloudflare - fonctionne parfaitement.
Données : nom, adresse, téléphone, site web, email

INSTALLATION (1 seule fois) :
  pip install playwright openpyxl
  playwright install chromium

LANCEMENT :
  python scrape_google_maps.py
"""

import asyncio
import re
import os
import csv
import random
from datetime import datetime

try:
    from playwright.async_api import async_playwright
except ImportError:
    print("❌ Lance d'abord : pip install playwright && playwright install chromium")
    exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_XLSX = os.path.join(OUTPUT_DIR, "avocats_paris.xlsx")
OUTPUT_CSV  = os.path.join(OUTPUT_DIR, "avocats_paris.csv")
TARGET = 500

# Recherches Google Maps à effectuer
RECHERCHES = [
    "cabinet avocat Paris 1er",
    "cabinet avocat Paris 2eme",
    "cabinet avocat Paris 3eme",
    "cabinet avocat Paris 4eme",
    "cabinet avocat Paris 5eme",
    "cabinet avocat Paris 6eme",
    "cabinet avocat Paris 7eme",
    "cabinet avocat Paris 8eme",
    "cabinet avocat Paris 9eme",
    "cabinet avocat Paris 10eme",
    "cabinet avocat Paris 11eme",
    "cabinet avocat Paris 12eme",
    "cabinet avocat Paris 13eme",
    "cabinet avocat Paris 14eme",
    "cabinet avocat Paris 15eme",
    "cabinet avocat Paris 16eme",
    "cabinet avocat Paris 17eme",
    "cabinet avocat Paris 18eme",
    "cabinet avocat Paris 19eme",
    "cabinet avocat Paris 20eme",
    "avocat droit des affaires Paris",
    "avocat droit du travail Paris",
    "avocat droit immobilier Paris",
    "avocat droit de la famille Paris",
    "avocat droit pénal Paris",
    "avocat fiscaliste Paris",
    "cabinet avocat Neuilly-sur-Seine",
    "cabinet avocat Boulogne-Billancourt",
    "cabinet avocat Levallois-Perret",
    "cabinet avocat Nanterre",
]

def extraire_email(texte):
    m = re.search(r"[\w.\-+]+@[\w.\-]+\.[a-zA-Z]{2,}", texte or "")
    return m.group(0).lower() if m else ""

async def pause(a=1, b=3):
    await asyncio.sleep(random.uniform(a, b))

async def scraper_fiche(page, url_fiche, nom):
    """Ouvre la fiche Google Maps d'un cabinet et extrait les détails."""
    cab = {"nom": nom, "adresse": "", "telephone": "", "email": "", "site_web": "", "zone": ""}
    try:
        await page.goto(url_fiche, wait_until="domcontentloaded", timeout=20000)
        await pause(2, 4)

        # Adresse
        for sel in ["[data-item-id='address'] .fontBodyMedium", "button[data-item-id='address']", "[aria-label*='Adresse']"]:
            el = page.locator(sel).first
            if await el.count() > 0:
                cab["adresse"] = (await el.inner_text()).strip()
                break

        # Téléphone
        for sel in ["[data-item-id*='phone'] .fontBodyMedium", "button[data-item-id*='phone']", "[aria-label*='Téléphone']", "[aria-label*='Phone']"]:
            el = page.locator(sel).first
            if await el.count() > 0:
                cab["telephone"] = (await el.inner_text()).strip()
                break

        # Site web
        for sel in ["a[data-item-id='authority']", "a[aria-label*='Site Web']", "a[href*='http'][data-item-id]"]:
            el = page.locator(sel).first
            if await el.count() > 0:
                href = await el.get_attribute("href")
                if href and "google" not in href and "maps" not in href:
                    cab["site_web"] = href
                    break

        # Email via le site web du cabinet
        if cab["site_web"]:
            try:
                await page.goto(cab["site_web"], wait_until="domcontentloaded", timeout=15000)
                await pause(1, 2)
                texte = await page.inner_text("body")
                cab["email"] = extraire_email(texte)
                if not cab["email"]:
                    # Chercher page contact
                    for sel in ["a[href*='contact']", "a:has-text('Contact')", "a:has-text('contact')"]:
                        lien = page.locator(sel).first
                        if await lien.count() > 0:
                            href_contact = await lien.get_attribute("href")
                            if href_contact:
                                if href_contact.startswith("/"):
                                    base = re.match(r"https?://[^/]+", cab["site_web"])
                                    href_contact = base.group(0) + href_contact if base else href_contact
                                await page.goto(href_contact, wait_until="domcontentloaded", timeout=10000)
                                await pause(1, 2)
                                texte = await page.inner_text("body")
                                cab["email"] = extraire_email(texte)
                            break
            except Exception:
                pass

    except Exception as e:
        pass

    return cab

async def scraper_liste(page, recherche, tous, noms_vus):
    """Scrape la liste Google Maps pour une recherche donnée."""
    url = f"https://www.google.com/maps/search/{recherche.replace(' ', '+')}"
    nouveaux = 0

    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=30000)
        await pause(3, 5)

        # Fermer popup consentement Google si présent
        for sel in ["button:has-text('Tout accepter')", "button:has-text('Accept all')", "#L2AGLb", "button:has-text('Accepter')"]:
            btn = page.locator(sel).first
            if await btn.count() > 0:
                await btn.click()
                await pause(1, 2)
                break

        # Défiler pour charger plus de résultats (Google Maps charge en scroll)
        panneau = page.locator("[role='feed'], .DxyBCb, .ecceSd").first
        for _ in range(8):
            if len(tous) >= TARGET:
                break
            try:
                await panneau.evaluate("el => el.scrollTop += 2000")
                await pause(1.5, 3)
            except Exception:
                await page.keyboard.press("End")
                await pause(1.5, 3)

        # Récupérer tous les résultats
        items = await page.query_selector_all("[role='feed'] > div, .Nv2PK, a.hfpxzc")

        fiches_urls = []
        noms_liste = []

        for item in items:
            try:
                # Nom
                nom_el = await item.query_selector(".fontHeadlineSmall, .qBF1Pd, h3, .NrDZNb, [jsan*='t-HiaYvf']")
                if not nom_el:
                    nom_el = await item.query_selector("a[aria-label]")
                    if nom_el:
                        nom = (await nom_el.get_attribute("aria-label") or "").strip()
                    else:
                        continue
                else:
                    nom = (await nom_el.inner_text()).strip()

                if not nom or nom.lower() in noms_vus:
                    continue
                if any(x in nom.lower() for x in ["google", "maps", "publicité"]):
                    continue

                # URL fiche
                lien = await item.query_selector("a[href*='maps/place'], a[href*='/maps/']")
                if not lien:
                    lien = await item.query_selector("a[aria-label]")
                href = await lien.get_attribute("href") if lien else ""

                if href and nom:
                    noms_vus.add(nom.lower())
                    fiches_urls.append(href)
                    noms_liste.append(nom)

            except Exception:
                continue

        print(f"    → {len(fiches_urls)} fiches trouvées")

        # Visiter chaque fiche pour les détails
        for nom, url_fiche in zip(noms_liste, fiches_urls):
            if len(tous) >= TARGET:
                break
            if nom.lower() in noms_vus and nom not in noms_liste:
                continue
            print(f"      Fiche : {nom[:50]}...", end="\r")
            cab = await scraper_fiche(page, url_fiche, nom)
            cab["zone"] = recherche
            tous.append(cab)
            nouveaux += 1
            await pause(2, 4)

    except Exception as e:
        print(f"    Erreur : {e}")

    return nouveaux

def exporter(donnees):
    if not donnees:
        print("⚠️  Aucune donnée.")
        return

    # CSV
    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        champs = ["nom", "adresse", "telephone", "email", "site_web", "zone"]
        w = csv.DictWriter(f, fieldnames=champs, extrasaction="ignore")
        w.writeheader()
        w.writerows(donnees)
    print(f"  CSV  → {OUTPUT_CSV}")

    if not EXCEL_OK:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Avocats Paris"
    BLEU = "1F3864"
    thin = Side(border_style="thin", color="CCCCCC")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"AVOCATS PARIS — {len(donnees)} contacts | {datetime.now().strftime('%d/%m/%Y')}"
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLEU)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    entetes = ["N°", "Nom / Cabinet", "Adresse", "Téléphone", "Email", "Site web", "Zone", "Statut"]
    for i, h in enumerate(entetes, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = brd
    ws.row_dimensions[2].height = 18
    ws.freeze_panes = "A3"

    for idx, d in enumerate(donnees, 1):
        r = idx + 2
        vals = [idx, d.get("nom",""), d.get("adresse",""), d.get("telephone",""),
                d.get("email",""), d.get("site_web",""), d.get("zone",""), "À contacter"]
        bg = "D6E4F0" if idx % 2 == 0 else "FFFFFF"
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = Font(size=10)
            c.border = brd
            c.alignment = Alignment(horizontal="center" if col in (1,4,8) else "left", vertical="center")
            if col == 5 and val:
                c.fill = PatternFill("solid", fgColor="E8F5E9")
            elif col == 5:
                c.fill = PatternFill("solid", fgColor="FFF3E0")
            else:
                c.fill = PatternFill("solid", fgColor=bg)
        ws.row_dimensions[r].height = 16

    for i, w in enumerate([4, 38, 40, 16, 34, 32, 30, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A2:H{len(donnees)+2}"
    wb.save(OUTPUT_XLSX)
    print(f"  XLSX → {OUTPUT_XLSX}")

async def main():
    print("\n" + "═"*55)
    print("  SCRAPER AVOCATS PARIS — GOOGLE MAPS")
    print(f"  Démarrage : {datetime.now().strftime('%H:%M:%S')}")
    print(f"  Objectif  : {TARGET} avocats")
    print("  (Une fenêtre Chrome va s'ouvrir — c'est normal)")
    print("═"*55)

    tous = []
    noms_vus = set()

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=False, slow_mo=50)
        context = await browser.new_context(
            viewport={"width": 1280, "height": 900},
            locale="fr-FR",
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        )
        page = await context.new_page()

        for i, recherche in enumerate(RECHERCHES):
            if len(tous) >= TARGET:
                break
            print(f"\n🔍 [{i+1}/{len(RECHERCHES)}] {recherche}")
            n = await scraper_liste(page, recherche, tous, noms_vus)
            print(f"    ✅ +{n} ajoutés | Total : {len(tous)}/{TARGET}")
            await pause(3, 6)

        await browser.close()

    emails = sum(1 for c in tous if c.get("email"))
    print(f"\n{'═'*55}")
    print(f"  RÉSULTAT : {len(tous)} cabinets")
    print(f"  Emails   : {emails} ({emails*100//max(len(tous),1)}%)")
    print("═"*55)

    exporter(tous)
    print(f"\n✅ Fichiers dans : {OUTPUT_DIR}")

asyncio.run(main())
