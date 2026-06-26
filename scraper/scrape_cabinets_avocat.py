"""
Scraper cabinets d'avocat Paris - Pages Jaunes
Lance : python scrape_cabinets_avocat.py
Résultat : cabinets_avocat_paris.xlsx dans le même dossier
"""

import requests
from bs4 import BeautifulSoup
import pandas as pd
import time
import random
import re
import os
from datetime import datetime

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_XLSX = os.path.join(OUTPUT_DIR, "cabinets_avocat_paris.xlsx")
TARGET = 500

ZONES = [
    "Paris 1er (75001)", "Paris 2eme (75002)", "Paris 3eme (75003)",
    "Paris 4eme (75004)", "Paris 5eme (75005)", "Paris 6eme (75006)",
    "Paris 7eme (75007)", "Paris 8eme (75008)", "Paris 9eme (75009)",
    "Paris 10eme (75010)", "Paris 11eme (75011)", "Paris 12eme (75012)",
    "Paris 13eme (75013)", "Paris 14eme (75014)", "Paris 15eme (75015)",
    "Paris 16eme (75016)", "Paris 17eme (75017)", "Paris 18eme (75018)",
    "Paris 19eme (75019)", "Paris 20eme (75020)",
    "Boulogne-Billancourt (92100)", "Neuilly-sur-Seine (92200)",
    "Levallois-Perret (92300)", "Nanterre (92000)", "Vincennes (94300)",
    "Saint-Denis (93200)", "Montreuil (93100)", "Versailles (78000)",
    "Courbevoie (92400)", "Issy-les-Moulineaux (92130)",
]

# Headers qui imitent un vrai navigateur Chrome
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8",
    "Accept-Language": "fr-FR,fr;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept-Encoding": "gzip, deflate, br",
    "Connection": "keep-alive",
    "Upgrade-Insecure-Requests": "1",
    "Sec-Fetch-Dest": "document",
    "Sec-Fetch-Mode": "navigate",
    "Sec-Fetch-Site": "none",
    "Sec-Fetch-User": "?1",
    "Cache-Control": "max-age=0",
}

def creer_session():
    """Crée une session avec cookies en visitant d'abord la page d'accueil."""
    session = requests.Session()
    session.headers.update(HEADERS)
    try:
        print("  Connexion à Pages Jaunes...")
        session.get("https://www.pagesjaunes.fr", timeout=15)
        time.sleep(random.uniform(2, 4))
    except Exception as e:
        print(f"  Avertissement : {e}")
    return session

def get_page(session, zone, page=1):
    url = "https://www.pagesjaunes.fr/annuaire/chercher"
    params = {"quoiqui": "avocat", "ou": zone, "page": page}
    try:
        resp = session.get(url, params=params, timeout=20)
        print(f"    Status {resp.status_code}")
        if resp.status_code == 200:
            return BeautifulSoup(resp.text, "lxml")
        elif resp.status_code == 403:
            print("    Bloqué - pause de 15 secondes...")
            time.sleep(15)
            # Recréer la session et réessayer
            session = creer_session()
            resp = session.get(url, params=params, timeout=20)
            if resp.status_code == 200:
                return BeautifulSoup(resp.text, "lxml")
        return None
    except Exception as e:
        print(f"    Erreur : {e}")
        return None

def extraire_email(texte):
    match = re.search(r"[\w.\-+]+@[\w.\-]+\.[a-zA-Z]{2,}", texte)
    return match.group(0) if match else ""

def parser_resultats(soup):
    cabinets = []
    cards = soup.select("li.bi-item, li.bi-pro, article.bi-item")
    if not cards:
        cards = soup.select("[class*='bi-item']")
    for card in cards:
        nom_tag = card.select_one("a.bi-denomination, h3 a, .denomination-links a, a[class*='denomination']")
        nom = nom_tag.get_text(strip=True) if nom_tag else ""
        if not nom:
            continue
        adr_tag = card.select_one("address, [class*='address']")
        adresse = adr_tag.get_text(" ", strip=True) if adr_tag else ""
        tel_tag = card.select_one("a[href^='tel:'], [class*='phone']")
        telephone = ""
        if tel_tag:
            telephone = tel_tag.get("href", "").replace("tel:", "") or tel_tag.get_text(strip=True)
        email = extraire_email(card.get_text())
        site_tag = card.select_one("a[class*='site'], a[class*='web']")
        site = site_tag.get("href", "") if site_tag else ""
        lien_tag = card.select_one("a.bi-denomination, h3 a")
        href = lien_tag.get("href", "") if lien_tag else ""
        fiche = "https://www.pagesjaunes.fr" + href if href.startswith("/") else ""
        cabinets.append({
            "nom": nom,
            "adresse": adresse,
            "telephone": telephone.strip(),
            "email": email,
            "site_web": site,
            "url_fiche": fiche,
        })
    return cabinets

def scraper_fiche(session, url):
    """Visite la fiche pour récupérer email et site web."""
    if not url:
        return {}
    try:
        time.sleep(random.uniform(1, 2))
        resp = session.get(url, timeout=15)
        if resp.status_code != 200:
            return {}
        soup = BeautifulSoup(resp.text, "lxml")
        email_tag = soup.select_one("a[href^='mailto:']")
        email = email_tag["href"].replace("mailto:", "").strip() if email_tag else extraire_email(soup.get_text())
        site_tag = soup.select_one("a[class*='site-web'], a[class*='website'], a[href*='http'][class*='site']")
        site = site_tag.get("href", "") if site_tag else ""
        return {"email": email, "site_web": site}
    except Exception:
        return {}

def exporter_excel(donnees):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cabinets Avocat Paris"

        BLEU = "1F3864"
        BLEU_CLAIR = "D6E4F0"
        thin = Side(border_style="thin", color="CCCCCC")
        bordure = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Titre
        ws.merge_cells("A1:I1")
        c = ws["A1"]
        c.value = f"CABINETS D'AVOCAT - PARIS MÉTROPOLE ({len(donnees)} contacts)"
        c.font = Font(bold=True, size=13, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        # En-têtes
        entetes = ["N°", "Nom", "Adresse", "Téléphone", "Email", "Site web", "Zone", "Statut", "Notes"]
        for i, h in enumerate(entetes, 1):
            c = ws.cell(row=2, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=BLEU)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = bordure
        ws.row_dimensions[2].height = 18
        ws.freeze_panes = "A3"

        # Données
        for idx, d in enumerate(donnees, 1):
            row = idx + 2
            vals = [
                idx, d.get("nom",""), d.get("adresse",""), d.get("telephone",""),
                d.get("email",""), d.get("site_web",""), d.get("zone",""), "À contacter", ""
            ]
            bg = BLEU_CLAIR if idx % 2 == 0 else "FFFFFF"
            for col, val in enumerate(vals, 1):
                c = ws.cell(row=row, column=col, value=val)
                c.font = Font(size=10)
                c.border = bordure
                c.alignment = Alignment(horizontal="center" if col in (1,4,8) else "left", vertical="center")
                if col == 5 and val:  # Email en vert
                    c.fill = PatternFill("solid", fgColor="E8F5E9")
                elif col == 5:        # Email vide en orange
                    c.fill = PatternFill("solid", fgColor="FFF3E0")
                else:
                    c.fill = PatternFill("solid", fgColor=bg)
            ws.row_dimensions[row].height = 16

        # Largeurs
        for i, w in enumerate([5, 38, 35, 16, 32, 32, 25, 14, 20], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        ws.auto_filter.ref = f"A2:I{len(donnees)+2}"
        wb.save(OUTPUT_XLSX)
        return True
    except Exception as e:
        print(f"Erreur Excel : {e}")
        # Fallback CSV
        pd.DataFrame(donnees).to_csv(OUTPUT_XLSX.replace(".xlsx", ".csv"), index=False, encoding="utf-8-sig")
        return False

# ─── PROGRAMME PRINCIPAL ──────────────────────────────────────────────────────

print("=" * 60)
print("  SCRAPER CABINETS D'AVOCAT - PARIS MÉTROPOLE")
print(f"  Démarrage : {datetime.now().strftime('%H:%M:%S')}")
print("=" * 60)

session = creer_session()
tous_cabinets = []
noms_vus = set()

for zone in ZONES:
    if len(tous_cabinets) >= TARGET:
        break

    print(f"\n📍 {zone}")
    soup = get_page(session, zone, page=1)
    if not soup:
        continue

    # Détecter nombre de pages
    nb_pages = 1
    pag = soup.select_one("[class*='pagination'] [class*='last'], .pagination li:last-child a")
    if pag:
        try:
            nb_pages = int(re.search(r"\d+", pag.get_text()).group())
        except Exception:
            nb_pages = 5

    for page in range(1, min(nb_pages + 1, 10)):
        if len(tous_cabinets) >= TARGET:
            break
        if page > 1:
            soup = get_page(session, zone, page)
            if not soup:
                break
            time.sleep(random.uniform(3, 6))

        nouveaux = 0
        for cab in parser_resultats(soup):
            cle = cab["nom"].lower().strip()
            if cle and cle not in noms_vus:
                noms_vus.add(cle)
                cab["zone"] = zone
                tous_cabinets.append(cab)
                nouveaux += 1

        print(f"    Page {page} : +{nouveaux} | Total : {len(tous_cabinets)}/{TARGET}")
        time.sleep(random.uniform(2, 5))

    time.sleep(random.uniform(4, 8))

# Enrichissement emails via fiches détaillées
print(f"\n📧 Recherche d'emails sur les fiches individuelles...")
enrichis = 0
for i, cab in enumerate(tous_cabinets):
    if cab.get("email") or not cab.get("url_fiche"):
        continue
    detail = scraper_fiche(session, cab["url_fiche"])
    if detail.get("email"):
        cab["email"] = detail["email"]
        enrichis += 1
    if detail.get("site_web") and not cab.get("site_web"):
        cab["site_web"] = detail["site_web"]
    if i % 50 == 0 and i > 0:
        print(f"    {i}/{len(tous_cabinets)} fiches visitées — {enrichis} emails trouvés")

# Export
print(f"\n✅ {len(tous_cabinets)} cabinets collectés — {enrichis} emails récupérés")
ok = exporter_excel(tous_cabinets)

print("\n" + "=" * 60)
if ok:
    print(f"  FICHIER EXCEL CRÉÉ : cabinets_avocat_paris.xlsx")
else:
    print(f"  FICHIER CSV CRÉÉ : cabinets_avocat_paris.csv")
print(f"  Dans le dossier : {OUTPUT_DIR}")
print("=" * 60)
