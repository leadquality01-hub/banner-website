"""
Scraper d'emails d'avocats parisiens individuels
Stratégie : Google → site perso de chaque avocat → email direct

INSTALLATION :
  pip install playwright openpyxl
  playwright install chromium

LANCEMENT :
  python scrape_emails_avocats.py

RÉSULTAT : emails_avocats_paris.xlsx
"""

import asyncio
import re
import os
import random
from datetime import datetime

try:
    from playwright.async_api import async_playwright
except ImportError:
    print("Lance d'abord :")
    print("  pip install playwright openpyxl")
    print("  playwright install chromium")
    exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_XLSX = os.path.join(OUTPUT_DIR, "emails_avocats_paris.xlsx")
OUTPUT_CSV  = os.path.join(OUTPUT_DIR, "emails_avocats_paris.csv")
TARGET = 500

# Requêtes Google pour trouver des avocats individuels avec leur site/email
REQUETES_GOOGLE = [
    "avocat Paris contact email @",
    "avocat droit des affaires Paris site perso contact",
    "avocat droit du travail Paris email contact",
    "avocat droit pénal Paris contact email",
    "avocat droit immobilier Paris email",
    "avocat droit de la famille Paris email contact",
    "avocat fiscaliste Paris email",
    "avocat divorce Paris email contact",
    "avocat Paris 75008 email contact",
    "avocat Paris 75016 email contact",
    "avocat Paris 75009 email contact",
    "avocat Paris 75017 email contact",
    "avocat Paris 75015 email contact",
    "avocat Paris 75006 email contact",
    "avocat Paris 75007 email contact",
    "avocat Paris 75001 email contact",
    "avocat contentieux Paris email",
    "avocat entreprise Paris email contact",
    "avocat succession Paris email",
    "avocat permis construire Paris email",
    "cabinet avocat Paris email site:*.fr",
    "maître avocat Paris email",
    "avocat associé Paris email contact",
    "avocat barreau paris email",
    "avocat conseil Paris email",
]

EMAIL_RE = re.compile(r"[\w.\-+]{2,}@[\w.\-]+\.[a-zA-Z]{2,6}")
DOMAINES_EXCLUS = {"example.com", "sentry.io", "google.com", "facebook.com",
                   "twitter.com", "youtube.com", "wix.com", "wordpress.com",
                   "w3.org", "schema.org", "adobe.com", "jquery.com"}

def est_email_valide(email):
    email = email.lower().strip()
    if len(email) < 6:
        return False
    domaine = email.split("@")[-1]
    if domaine in DOMAINES_EXCLUS:
        return False
    # Garder emails avec domaines cabinet/perso
    return True

async def pause(a=1, b=3):
    await asyncio.sleep(random.uniform(a, b))

async def extraire_emails_page(page):
    """Extrait tous les emails d'une page web."""
    try:
        texte = await page.inner_text("body")
        html  = await page.content()
        # Emails dans le texte visible
        emails_texte = EMAIL_RE.findall(texte)
        # Emails dans les liens mailto:
        emails_mailto = await page.evaluate("""
            () => [...document.querySelectorAll('a[href^="mailto:"]')]
                  .map(a => a.href.replace('mailto:','').split('?')[0].trim())
        """)
        tous = set(emails_texte + emails_mailto)
        return [e.lower() for e in tous if est_email_valide(e)]
    except Exception:
        return []

async def visiter_page_contact(page, base_url):
    """Cherche et visite la page Contact du site pour trouver l'email."""
    try:
        for sel in [
            "a:has-text('Contact')", "a:has-text('contact')",
            "a:has-text('Nous contacter')", "a:has-text('Contactez')",
            "a[href*='contact']", "a[href*='Contact']",
        ]:
            lien = page.locator(sel).first
            if await lien.count() > 0:
                href = await lien.get_attribute("href")
                if href:
                    if href.startswith("/"):
                        href = base_url.rstrip("/") + href
                    elif not href.startswith("http"):
                        href = base_url.rstrip("/") + "/" + href
                    await page.goto(href, wait_until="domcontentloaded", timeout=12000)
                    await pause(1, 2)
                    return await extraire_emails_page(page)
    except Exception:
        pass
    return []

async def scraper_site_avocat(page, url):
    """Visite le site d'un avocat et extrait son email."""
    try:
        await page.goto(url, wait_until="domcontentloaded", timeout=15000)
        await pause(1, 2)

        base_url = re.match(r"https?://[^/]+", url)
        base = base_url.group(0) if base_url else url

        # Chercher emails sur la page d'accueil
        emails = await extraire_emails_page(page)

        # Si pas trouvé → aller sur la page Contact
        if not emails:
            emails = await visiter_page_contact(page, base)

        return emails
    except Exception:
        return []

async def scraper_google(page, requete, resultats_vus):
    """Cherche sur Google et retourne les URLs de sites d'avocats."""
    urls = []
    try:
        url_recherche = f"https://www.google.fr/search?q={requete.replace(' ', '+')}&num=20&hl=fr"
        await page.goto(url_recherche, wait_until="domcontentloaded", timeout=20000)
        await pause(2, 4)

        # Accepter cookies Google si popup
        for sel in ["button:has-text('Tout accepter')", "#L2AGLb", "button:has-text('Accept all')"]:
            btn = page.locator(sel).first
            if await btn.count() > 0:
                await btn.click()
                await pause(1, 2)
                break

        # Extraire les URLs des résultats
        liens = await page.evaluate("""
            () => [...document.querySelectorAll('a[href]')]
                  .map(a => a.href)
                  .filter(h => h.startsWith('http') && !h.includes('google') && !h.includes('youtube'))
        """)

        for lien in liens:
            if lien not in resultats_vus:
                # Filtrer pour garder seulement les sites qui ressemblent à des sites d'avocats
                if any(k in lien.lower() for k in ['avocat', 'juriste', 'cabinet', 'law', 'legal', 'maitre', 'barreau', 'conseil']):
                    urls.append(lien)
                    resultats_vus.add(lien)

    except Exception as e:
        print(f"    Erreur Google : {e}")

    return urls[:10]  # Max 10 URLs par recherche

async def main():
    print("\n" + "═"*55)
    print("  EMAILS D'AVOCATS PARISIENS — SCRAPER")
    print(f"  Démarrage : {datetime.now().strftime('%H:%M:%S')}")
    print(f"  Objectif  : {TARGET} emails")
    print("═"*55)

    contacts = []      # [{nom, email, site, specialite}]
    emails_vus = set()
    urls_vus   = set()

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=False, slow_mo=80)
        context = await browser.new_context(
            viewport={"width": 1280, "height": 900},
            locale="fr-FR",
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        )
        page = await context.new_page()

        for i, requete in enumerate(REQUETES_GOOGLE):
            if len(contacts) >= TARGET:
                break

            print(f"\n🔍 [{i+1}/{len(REQUETES_GOOGLE)}] {requete}")
            urls = await scraper_google(page, requete, urls_vus)
            print(f"    → {len(urls)} sites trouvés")

            for url in urls:
                if len(contacts) >= TARGET:
                    break
                try:
                    print(f"    Visite : {url[:60]}...", end="\r")
                    emails_trouves = await scraper_site_avocat(page, url)

                    for email in emails_trouves:
                        if email not in emails_vus and est_email_valide(email):
                            emails_vus.add(email)
                            # Essayer d'extraire le nom depuis le titre de la page
                            try:
                                titre = await page.title()
                            except Exception:
                                titre = ""
                            contacts.append({
                                "email": email,
                                "nom": titre[:60].strip(),
                                "site_web": url,
                                "specialite": extraire_specialite(requete),
                                "telephone": "",
                            })
                    if emails_trouves:
                        print(f"    ✅ {url[:50]} → {len(emails_trouves)} email(s) | Total : {len(contacts)}")

                except Exception:
                    pass
                await pause(1.5, 3)

            await pause(3, 6)

        await browser.close()

    print(f"\n{'═'*55}")
    print(f"  RÉSULTAT : {len(contacts)} emails collectés")
    print("═"*55)

    exporter(contacts)
    print(f"\n✅ Fichier : {OUTPUT_XLSX or OUTPUT_CSV}")

def extraire_specialite(requete):
    mots = ["droit des affaires", "droit du travail", "droit pénal", "droit immobilier",
            "droit de la famille", "fiscaliste", "divorce", "contentieux", "succession"]
    for m in mots:
        if m in requete.lower():
            return m.title()
    return "Droit général"

def exporter(donnees):
    import csv
    if not donnees:
        print("⚠️  Aucun email trouvé.")
        return

    # CSV
    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["email","nom","site_web","specialite","telephone"], extrasaction="ignore")
        w.writeheader()
        w.writerows(donnees)
    print(f"  CSV  → {OUTPUT_CSV}")

    if not EXCEL_OK:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Emails Avocats Paris"
    BLEU = "1F3864"
    thin = Side(border_style="thin", color="CCCCCC")
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"EMAILS AVOCATS PARIS — {len(donnees)} contacts | {datetime.now().strftime('%d/%m/%Y')}"
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLEU)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    entetes = ["N°", "EMAIL", "Nom / Cabinet", "Spécialité", "Site web", "Statut"]
    for i, h in enumerate(entetes, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = brd
    ws.row_dimensions[2].height = 20
    ws.freeze_panes = "A3"

    for idx, d in enumerate(donnees, 1):
        r = idx + 2
        vals = [idx, d.get("email",""), d.get("nom",""), d.get("specialite",""),
                d.get("site_web",""), "À contacter"]
        bg = "D6E4F0" if idx % 2 == 0 else "FFFFFF"
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = Font(size=10)
            c.border = brd
            c.alignment = Alignment(horizontal="center" if col in (1,6) else "left", vertical="center")
            # Email en vert bien visible
            if col == 2:
                c.fill = PatternFill("solid", fgColor="E8F5E9")
                c.font = Font(size=10, bold=True, color="1B5E20")
            else:
                c.fill = PatternFill("solid", fgColor=bg)
        ws.row_dimensions[r].height = 16

    for i, w in enumerate([4, 38, 40, 24, 40, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A2:F{len(donnees)+2}"
    wb.save(OUTPUT_XLSX)
    print(f"  XLSX → {OUTPUT_XLSX}")

asyncio.run(main())
